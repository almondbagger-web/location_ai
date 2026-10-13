"""道路使用・牽引・公園許可の社内表を permit_knowledge.json に変換する。"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import xlrd
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent
ROAD_XLS = Path("/Users/user/Desktop/個人諸々/ALMONDBAGGER/ホームページ/東京都道路使用申請日数 2012.7.20.xls")
PARK_XLSX = Path("/Users/user/Desktop/番組別/撮影関係/挨拶回り・道路使用フォーマット/(東京都)公園撮影許可一覧.xlsx")
POLICE_PDF = Path("/Users/user/Desktop/番組別/撮影関係/挨拶回り・道路使用フォーマット/道路使用　警察署一覧【東京都】.pdf")
OUT = BASE / "permit_knowledge.json"

HEADER_WORDS = {"止め撮り", "走行撮影", "牽引", "中日", "その他", "区", "地区", "警察署名"}


def clean(value) -> str:
    text = str(value or "")
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def station_key(name: str) -> str:
    name = clean(name).replace("䇮原", "荏原").replace("四ツ谷", "四谷").replace("四ッ谷", "四谷")
    name = name.replace("【", "").replace("】", "")
    return name.replace("警察署", "")


def mark(token: str) -> str:
    token = clean(token)
    if token in ("×", "✕", "✖"):
        return "不可"
    if token in ("〇", "○", "◯"):
        return "可"
    if token in ("△", "▲"):
        return "応相談"
    if "応相談" in token:
        return "応相談"
    if "客観" in token:
        return "客観のみ可"
    return token


def load_road_xls() -> list[dict]:
    book = xlrd.open_workbook(ROAD_XLS)
    stations = []
    for sheet in book.sheets():
        area = ""
        for row_index in range(sheet.nrows):
            values = [clean(sheet.cell_value(row_index, col)) for col in range(min(sheet.ncols, 8))]
            if not any(values):
                continue
            if values[0] in ("地区名", "23区", "東京23区", "東京都23区外"):
                continue
            if values[0]:
                area = values[0]
            name = values[1]
            if "警察" not in name:
                continue
            stations.append(
                {
                    "area": area,
                    "name": name.replace("四ツ谷", "四谷"),
                    "phone": values[2],
                    "address": values[3],
                    "lead_days_2012": values[4],
                    "period": values[5],
                    "note": " ".join(part for part in values[6:] if part),
                }
            )
    return stations


def page_items(page) -> list[dict]:
    items = []

    def visitor(text, cm, tm, font_dict, font_size):
        token = clean(text)
        if token:
            items.append({"y": tm[5], "x": tm[4], "t": token})

    page.extract_text(visitor_text=visitor)
    return items


def load_police_pdf() -> dict[str, dict]:
    reader = PdfReader(str(POLICE_PDF))
    found: dict[str, dict] = {}
    for page in reader.pages:
        items = page_items(page)
        stations = [item for item in items if "警察署】" in item["t"] or item["t"].endswith("警察署】")]
        for station in stations:
            near = [
                item
                for item in items
                if abs(item["y"] - station["y"]) <= 14 and 180 < item["x"] < 760 and item["t"] not in HEADER_WORDS
            ]
            still = drive = tow = lead = ""
            jurisdiction = []
            for item in sorted(near, key=lambda row: row["x"]):
                token = item["t"]
                is_mark = token in ("×", "〇", "○", "◯", "△") or "応相談" in token or "客観" in token
                if item["x"] >= 370:
                    if item["x"] < 740:
                        jurisdiction.append(token)
                    continue
                if re.fullmatch(r"\d日", token):
                    lead = token
                    continue
                if not is_mark:
                    continue
                if item["x"] < 240:
                    if not still:
                        still = mark(token)
                elif item["x"] < 300:
                    if not drive:
                        drive = mark(token)
                elif item["x"] < 345:
                    if not tow:
                        tow = mark(token)
            notes = [
                item["t"]
                for item in items
                if item["x"] > 740 and abs(item["y"] - station["y"]) <= 18 and item["t"] not in HEADER_WORDS
            ]
            key = station_key(station["t"])
            areas = []
            seen = set()
            for piece in jurisdiction:
                piece = piece.strip("、 ")
                if piece and piece not in seen and piece not in ("：", "・"):
                    seen.add(piece)
                    areas.append(piece)
            found[key] = {
                "still": still,
                "drive": drive,
                "tow": tow,
                "lead_days_2018": lead,
                "jurisdiction": clean(" ".join(areas))[:500],
                "pdf_note": clean(" ".join(notes))[:240],
            }
    return found


def load_parks() -> list[dict]:
    book = openpyxl.load_workbook(PARK_XLSX, data_only=True)
    sheet = book[book.sheetnames[0]]
    parks = []
    for row in sheet.iter_rows(values_only=True):
        area = clean(row[0] if row else "")
        if not area or area in ("23区", "２３区外", "市町村", "東京都建設局"):
            continue
        contact = clean(row[2] if len(row) > 2 else "")
        if area in ("撮影使用",) or contact in ("問い合わせ窓口", "公園管理課"):
            continue
        updated = row[12] if len(row) > 12 else ""
        if isinstance(updated, datetime):
            updated = updated.strftime("%Y.%m.%d")
        parks.append(
            {
                "area": area,
                "allowed": clean(row[1] if len(row) > 1 else ""),
                "contact": contact,
                "days": clean(row[3] if len(row) > 3 else ""),
                "hours": clean(row[4] if len(row) > 4 else ""),
                "night": clean(row[5] if len(row) > 5 else ""),
                "time_limit": clean(row[6] if len(row) > 6 else ""),
                "people": clean(row[7] if len(row) > 7 else ""),
                "ng": clean(row[8] if len(row) > 8 else ""),
                "fee": clean(row[9] if len(row) > 9 else ""),
                "note": clean(row[10] if len(row) > 10 else ""),
                "updated": clean(updated),
            }
        )
    return parks


def main():
    stations = load_road_xls()
    pdf_rows = load_police_pdf()
    used = set()
    for station in stations:
        extra = pdf_rows.get(station_key(station["name"]), {})
        used.add(station_key(station["name"]))
        station.update(
            {
                "still": extra.get("still", ""),
                "drive": extra.get("drive", ""),
                "tow": extra.get("tow", ""),
                "lead_days_2018": extra.get("lead_days_2018", ""),
                "jurisdiction": extra.get("jurisdiction", ""),
                "pdf_note": extra.get("pdf_note", ""),
            }
        )
    for key, extra in pdf_rows.items():
        if key in used:
            continue
        stations.append(
            {
                "area": "",
                "name": f"{key}警察署",
                "phone": "",
                "address": "",
                "lead_days_2012": "",
                "period": "",
                "note": "",
                **extra,
            }
        )
    parks = load_parks()
    payload = {
        "road_as_of": "2012.7.20",
        "tow_as_of": "2018年12月",
        "park_as_of": "2022.6.22",
        "stations": stations,
        "parks": parks,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    tow_ok = [row["name"] for row in stations if row.get("tow") == "可"]
    tow_ask = [row["name"] for row in stations if row.get("tow") == "応相談"]
    missing_tow = [row["name"] for row in stations if not row.get("tow")]
    print(f"stations {len(stations)} parks {len(parks)}")
    print("牽引可", len(tow_ok), " ".join(tow_ok))
    print("牽引応相談", " ".join(tow_ask))
    print("牽引欄なし", len(missing_tow), " ".join(missing_tow[:20]))


if __name__ == "__main__":
    main()
